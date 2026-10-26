"""Generate the local training draft and its companion from a single curriculum.
Run with the bundled document Python runtime. No platform writes are performed.
"""
import json
import base64
import html
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'course-creation'
OUT.mkdir(parents=True, exist_ok=True)

def text(content, style='plain'):
    return dict(content_type='text', text_type=style, content=content, is_published=False)

def quiz(kind, question, answers, explanation, options=None):
    return dict(quiz_type=kind, question=question, options=options or [], correct_answer=answers,
                explanation=explanation, image_url=None, time_limit_seconds=None, is_published=False)

def lesson(title, objective, paragraphs, task, quizzes=None, blocks=None, kind='reading'):
    return dict(title=title, description=objective, lesson_type=kind,
                learning_objectives=[objective], estimated_duration_minutes=10,
                mini_lessons=[dict(title=title, description=objective, is_published=False,
                    blocks=[*[text(p) for p in paragraphs], *(blocks or []), text('Try it yourself\n'+task, 'remember')])],
                quizzes=quizzes or [])

modules = [
    dict(title='Plan a course learners can complete', description='Define the audience and a useful outcome.', lessons=[
        lesson('Choose the audience and outcome', 'Write an observable outcome and a small course outline.', [
            'Start with what an educator or tutor should be able to do after the course. For this manual, the outcome is to create, preview and submit a complete UniVybe course with appropriate content and quizzes.',
            'Write outcomes using actions you can observe: create, explain, solve or demonstrate. For example: Create one module containing a lesson, a mini-lesson, content blocks and a quiz. Avoid an outcome such as Understand everything about teaching.',
            'Identify what learners already know, what they need to bring, and which device they will use. Keep each lesson focused on one outcome. Estimate the time only after drafting and walking through the activities.'
        ], 'Write a course title, target audience, two learning objectives and one prerequisite. Sketch two modules with one focused lesson in each.'),
        lesson('Understand the UniVybe course structure', 'Put each item in the correct place in the course outline.', [
            'A course is the full learning experience. Modules divide it into topics. Lessons sit inside modules. Mini-lessons divide each lesson into focused screens. Text and media blocks sit inside a mini-lesson.',
            'Quizzes belong directly to a lesson, alongside its mini-lessons. They do not belong inside a content block. The student player builds screens from mini-lessons and quizzes.',
            'The Academic module field in course details links an official university registry entry. It is different from the modules you create as chapters in this course.'
        ], 'Draw your course tree and mark where a quiz belongs.', [
            quiz('mcq', 'Where do you add a quiz?', ['To a lesson'], 'A lesson holds mini-lessons and quizzes.', ['To a lesson', 'Inside a content block', 'To the course thumbnail'])
        ], [text('Course → Module → Lesson → Mini-lesson → Content blocks\nLesson → Quizzes', 'simply_put')])
    ]),
    dict(title='Create the draft and build the outline', description='Create a small complete learning path before adding more material.', lessons=[
        lesson('Create course details', 'Save a course draft with clear learner facing information.', [
            'Open Courses and create a course. Enter the title and description, then save. These are required in the current frontend. State who the course is for, what learners will do and what they need before starting. Use the short description as a concise summary.',
            'Choose General for a broad audience, or the appropriate academic level. Not set leaves the level unspecified. Add university, faculty, department, organization or an official academic module only when relevant; an organization course does not require a university context.',
            'Add learning objectives, prerequisites and tags using one item per line. Add a cover URL or use the upload controls where available. Check that the image is readable on a small screen and that you have permission to use it.',
            'Estimated hours and lesson or quiz counts describe the intended scope. Check them against the final outline. Pricing and coin rewards entered by creators are suggestions; an admin sets final pricing and rewards.'
        ], 'Create a practice draft with a specific title, a useful description and two learning objectives. Save it and confirm that it appears in Courses.'),
        lesson('Build modules lessons and mini lessons', 'Save and reorder a module containing one complete lesson.', [
            'In the course outline, choose Add module. Give it a topic title and save. Expand that module, choose Add lesson, set its title and lesson type, and save. Reading, video and mixed describe the lesson; media itself is added as content blocks.',
            'Expand the lesson and choose Add mini-lesson. Give the screen a meaningful title and save. Expand it and choose Add content block. Add one plain text block first, save, then preview the course.',
            'Each editor saves separately. Save before choosing a different outline item. Use the saved confirmation to distinguish a confirmed save from a retained unsaved draft. Moving away with unsaved changes opens a save or discard decision.',
            'Use drag handles or Move up and Move down to reorder siblings. Order applies within the selected parent. After an interrupted reorder, reload and inspect the saved sequence before trying again.'
        ], 'Build one module with a reading lesson and two mini-lessons. Add a text block to each, reorder them and verify the result in preview.', [
            quiz('one_word', 'What holds a lesson in the course outline? Answer with one singular word.', ['module'], 'A module contains lessons.')
        ])
    ]),
    dict(title='Use every supported content type', description='Choose a format that supports the learning task.', lessons=[
        lesson('Write text explanations and reminders', 'Use the four explanation styles with a clear purpose.', [
            'Choose Content type text, then choose the text style. Plain is the default for explanations, instructions and worked examples. Break long passages into focused mini-lessons instead of placing the whole course in one block.',
            'Simply put gives a simpler restatement after an explanation. Important highlights a consequence or condition. Remember reinforces something learners need to carry into the next activity. Keep these blocks brief so their emphasis remains useful.'
        ], 'Teach one idea using plain text, then add a simpler restatement, one important condition and one reminder.', blocks=[
            text('A learning objective describes something the learner can demonstrate.', 'plain'),
            text('Write what the learner will do.', 'simply_put'),
            text('Only add media you are allowed to share with learners.', 'important'),
            text('Save each content block before moving to another item.', 'remember')]),
        lesson('Add formulas links and copyable prompts', 'Create a formula a labelled link and a reusable prompt.', [
            'For Formula, enter LaTeX and inspect the formula preview. For example, \\frac{3}{4} teaches a fraction. Explain the symbols in a nearby plain text block so the formula is meaningful without relying only on the visual.',
            'For Link, enter one complete HTTP or HTTPS address as the content. The optional label is limited to 100 characters on one line and must not contain a web address. Use Check link to confirm the destination. Put longer instructions in a preceding plain text block.',
            'For Prompt, enter plain text learners can copy. Preserve useful line breaks and placeholders. State what to replace before using it, and check the student Copy prompt control. A prompt is a content format rather than a graded quiz.'
        ], 'Create a formula, a labelled course reference link and a prompt with a replaceable topic. Check their previews.', blocks=[
            text(r'\frac{3}{4}', 'formula'),
            text('Create three practice questions about [TOPIC].\nUse beginner level language.\nExplain each answer after I attempt it.', 'prompt')]),
        lesson('Add images video embeds and documents', 'Select and check each of the four media block types.', [
            'Choose Image for an illustration or screenshot. Add a valid media URL or use the lesson media upload controls. Add a caption using advanced metadata when needed. Explain essential visual information in an adjacent text block because captions are not a substitute for verified accessible alternatives.',
            'Choose Video for a demonstration. Paste the media address, including a supported YouTube watch, share or embed URL, rather than iframe code. Provide a short text summary and an accessible transcript or captions, then check playback on desktop and mobile.',
            'Choose Embed for an external interactive resource. Check that the provider allows embedding and that learners can access it without your account. Provide a link or written alternative if embedding is blocked.',
            'Choose Document for a handout or reference file. Explain what learners should do with the file and test opening it as a learner. A document block displays or links media; course level resource attachments are managed separately.',
            'The media URL is required for non-text blocks. Advanced media metadata accepts JSON. Leave existing metadata unchanged unless needed; do not erase unfamiliar saved properties. For a caption, an example is {"caption":"Course outline example"}.'
        ], 'Add one image, one short video, one allowed embed and one document from your own approved materials. Test every address and include a text alternative.', kind='mixed')
    ]),
    dict(title='Build useful quizzes', description='Create examples of every authorable quiz type.', lessons=[
        lesson('Write multiple choice questions', 'Create questions with one or several correct choices.', [
            'Expand a lesson and choose Add quiz. Choose Multiple choice, write a clear question, add at least two distinct choices and tick every correct choice. For a single correct answer, tick only one. For multiple correct answers, say Select all that apply in the question.',
            'Use plausible incorrect choices that reveal a misunderstanding. Avoid trick wording or choices that differ only in punctuation. Add an explanation describing why the answer is correct and what the learner should revisit.',
            'Open optional settings only when needed. A blank time limit means untimed. Empty reward and attempt fields keep existing values or use server defaults for new quizzes. These defaults need to be confirmed with the backend before making learner promises.'
        ], 'Create one question with one correct choice and another with two correct choices. Inspect the answer key in preview.', [
            quiz('mcq', 'Select all items that belong directly to a lesson.', ['Mini-lessons', 'Quizzes'], 'Mini-lessons and quizzes belong to a lesson; blocks belong to mini-lessons.', ['Mini-lessons', 'Quizzes', 'Course thumbnails'])
        ]),
        lesson('Write short answer and one word questions', 'Choose accepted answers that fit the intended response.', [
            'Short answer accepts a short response and alternative wording. Add the main answer first and use Add alternative answer for valid variations. Keep the prompt specific enough that a short answer can demonstrate the outcome.',
            'One word requires each accepted answer to be a single word. Use alternative answer rows for legitimate spelling or terminology variations. Do not add a sentence to this type.',
            'These controls do not provide a teacher marked essay workflow. Avoid asking for a long explanation and expecting automatic grading. Check correct and incorrect answers in the student portal, including capitalization and spacing, before release.'
        ], 'Write one short answer question and one one-word question. Enter a valid alternative answer for the short answer question.', [
            quiz('short_answer', 'What should you do before leaving an edited block?', ['Save the block', 'Save changes'], 'Save the current editor before moving on.'),
            quiz('one_word', 'Which text style displays mathematical LaTeX?', ['formula'], 'Use the formula text style for LaTeX.')
        ]),
        lesson('Write fill in the blank and tap to fill questions', 'Keep blank answers and word bank choices in the intended order.', [
            'For Fill in the blank, write the sentence with blank spaces and enter answers in left to right blank order. Use Add blank answer for another position. Move the answer rows up or down to fix their order.',
            'For Tap to fill, add the word bank first, then select the answer for each blank from that bank. Include useful distractors. Every selected answer must match a word bank entry.',
            'The tutor form does not define a reliable blank marker or count blanks for you. Spaces at the start and end of the question are trimmed by the server. Keep blanks between visible text and verify their display and grading in the student portal. The draft includes both types for review, but their end to end release is gated on the compatibility checks in the review log.'
        ], 'Create one two-blank question and one tap-to-fill question. Check answer order, then try correct and incorrect submissions in a controlled student session.', [
            quiz('fill_blank', 'A course contains  and a module contains .', ['modules', 'lessons'], 'The hierarchy is course, module, lesson, mini-lesson, block.'),
            quiz('tap_fill', 'Add content blocks to a  and add quizzes to a .', ['mini-lesson', 'lesson'], 'Blocks belong to mini-lessons. Quizzes belong to lessons.', ['mini-lesson', 'lesson', 'thumbnail'])
        ])
    ]),
    dict(title='Review submit and maintain the course', description='Inspect the learner experience and record improvements.', lessons=[
        lesson('Review the student experience', 'Complete a structured desktop and mobile review.', [
            'Use Preview course to inspect the saved content and answer review. This preview helps with authoring; it does not prove enrollment, grading, progress, media access or publication behavior in the student portal.',
            'In a controlled student session, inspect the course summary, open the first lesson and read every mini-lesson. Test links, copy prompts, formulas and media on desktop and a small screen. Check text readability and whether the lesson order makes sense.',
            'Try each quiz with correct, incorrect and alternative answers. Confirm the explanation, allowed attempts, next screen behavior and recorded progress. Check lesson completion and access to the next lesson. Do not treat the screenshots in the companion as proof of successful grading; they use illustrative data.',
            'For each difficulty, record the exact action, what you expected, what happened and a suggested change. Separate missing content from an editor problem or student compatibility issue. Retest after each change.'
        ], 'Walk through your practice course on desktop and mobile. Record at least one improvement with steps to reproduce and a clear expected result.'),
        lesson('Submit a complete course for approval', 'Prepare a course for review and retain a maintenance checklist.', [
            'Check the title, description, objectives, prerequisites, counts and estimated time. Confirm every module, mini-lesson, block and quiz that should be visible is marked available where the form provides that setting. Saving a draft and enabling an item are separate from course approval.',
            'Save the current editor before choosing Submit for approval. The creation portal submits the course for review and expects a pending approval status. An admin decision is required; submitting does not by itself make the course public.',
            'After a review decision, inspect the returned status and feedback. Apply the requested corrections and follow the available submission workflow. If a save result is uncertain, use the check saved controls, inspect the server version and avoid repeating a create request that might duplicate content.',
            'Before launch, resolve the quiz compatibility checks, replace all pending media and screenshots, and complete the controlled student walkthrough. Keep a dated record of content changes, broken link repairs and learner feedback.'
        ], 'Prepare your course checklist and submit only when the actual course is complete. This local training draft remains unpublished.')
    ])
]

# The draft follows persisted field names and nesting, but is deliberately not
# represented as an API payload. Assets without final URLs are tracked separately.
for mi, module in enumerate(modules):
    module.update(order_index=mi, is_published=False)
    for li, item in enumerate(module['lessons']):
        item['order_index'] = li
        for ni, mini in enumerate(item['mini_lessons']):
            mini['order_index'] = ni
            for bi, block in enumerate(mini['blocks']): block['order_index'] = bi
        for qi, q in enumerate(item['quizzes']): q['order_index'] = qi

assets = [
    dict(key='reference_link', content_type='text', text_type='link', target='module 3 lesson 2', instruction='Use the final published companion document URL with label Course creation companion.'),
    dict(key='course_outline_image', content_type='image', target='module 3 lesson 3', instruction='Upload screenshots/02-student-lesson.png and explain the course outline.'),
    dict(key='course_creation_video', content_type='video', target='module 3 lesson 3', instruction='Record the actual educator workflow with captions and a transcript.'),
    dict(key='interactive_embed', content_type='embed', target='module 3 lesson 3', instruction='Choose an approved embeddable learning activity and check learner access.'),
    dict(key='companion_document', content_type='document', target='module 3 lesson 3', instruction='Upload course-creation-manual.docx or an approved exported reading copy and use the resulting URL.')
]
course = dict(title='Create your first UniVybe course', short_description='A practical course creation manual for educators and tutors.', description='Plan, build, preview and prepare a course for approval using UniVybe modules, lessons, mini-lessons, content blocks and quizzes.', level='general', learning_objectives=['Create a complete course outline', 'Select each supported text and media format', 'Author all five quiz types', 'Check the learner experience and prepare a course for review'], prerequisites=['Access to an approved educator or tutor account', 'A topic to teach and permission to use its learning materials'], tags=['educator training', 'course creation'], suggested_is_paid=False, estimated_duration_hours=3, estimated_lessons_count=sum(len(m['lessons']) for m in modules), estimated_quizzes_count=sum(len(l['quizzes']) for m in modules for l in m['lessons']))
draft = dict(format='univybe-local-course-draft', version=1, updated='2026-10-26', state='local draft requiring assets and integration review', course=course, modules=modules, pending_assets=assets, note='No bulk course import contract is implemented. Create each parent before its children using the current editor or a separately reviewed importer. Use returned parent IDs, never array positions. All authorable publication flags remain false.')
(OUT/'course-draft.json').write_text(json.dumps(draft, indent=2, ensure_ascii=False)+'\n', encoding='utf-8')

lines = ['# UniVybe course creation manual', '', 'Local draft for educators and tutors · 26 October 2026', '', 'Use this manual to plan, build, check and prepare a UniVybe course for approval. The companion training draft uses the same course hierarchy and provides practice tasks and worked quiz examples. It has not been imported or published.', '', '## How to use this manual', '', 'Work through the modules in order and build a small practice course as you go. Complete the task at the end of each lesson. Use the student screenshots to understand the learner view, then verify the real course in a controlled student session.', '', 'The screenshots were captured from the existing Edu_Student_g1 student build on 26 October 2026 with illustrative API responses. They show actual interface screens, not a published training course or verified grading results.', '', '## Student portal walkthrough', '', 'Open the course overview to check the title, description and curriculum. Open a lesson to see mini-lessons and quizzes as successive screens. The outline shows the learning sequence; Next advances through the current lesson. Completing a lesson and saving progress are separate from simply viewing its screens.', '']
for filename, caption in [('01-student-course.png','Student course overview with illustrative training content'), ('02-student-lesson.png','Student lesson showing plain text Simply put and Remember blocks'), ('03-student-quiz.png','Student quiz with one correct choice before submission'), ('04-student-mobile-quiz.png','The same quiz on a narrow mobile screen')]:
    lines += [f'![{caption}](screenshots/{filename})', '', caption+'. Illustrative data.', '']
for mi, module in enumerate(modules, 1):
    lines += [f'## Module {mi} {module["title"]}', '', module['description'], '']
    for li, item in enumerate(module['lessons'], 1):
        lines += [f'### Lesson {mi}.{li} {item["title"]}', '', 'Learning outcome: '+item['description'], '']
        for block in item['mini_lessons'][0]['blocks']:
            prefix = '' if block['text_type']=='plain' else block['text_type'].replace('_',' ').capitalize()+': '
            lines += [prefix+block['content'], '']
        for q in item['quizzes']:
            lines += ['Quiz example — '+q['quiz_type'].replace('_',' '), '', q['question'], '']
            if q['options']: lines += ['Choices: '+' | '.join(q['options']), '']
            lines += ['Answer key: '+' | '.join(q['correct_answer']), '', 'Explanation: '+q['explanation'], '']
lines += ['## Coverage and assets before release', '', 'All five authorable quiz types are included: multiple choice, short answer, one word, fill in the blank and tap to fill. Multiple choice includes a single answer example and a multiple answer example. True or false can be written with two choices; the tutor editor does not offer a separate true or false or teacher marked essay type.', '', 'All five content types are taught: text, image, video, embed and document. All seven text styles are taught: plain, simply put, important, remember, formula, link and prompt. Link and media examples need final approved URLs before they can become valid saved blocks. They are listed as pending assets instead of fake working URLs.', '']
for asset in assets: lines += ['- '+asset['instruction'], '']
lines += ['## Final course checklist', '']
for s in ['The audience objectives and prerequisites are clear.', 'The saved outline matches the intended module and lesson order.', 'All intended visible items have their availability settings checked.', 'Every text style media type and quiz type has been reviewed.', 'Media and external links work without educator account access.', 'Desktop and mobile reading and navigation have been checked.', 'Correct incorrect and alternative quiz answers have been tested against the real API.', 'Completion progress and next lesson access have been checked.', 'The review log has no unresolved release blocking issues.', 'The course has been submitted for approval only after all editors are saved.']: lines += ['- [ ] '+s]
(OUT/'course-creation-manual.md').write_text('\n'.join(lines)+'\n', encoding='utf-8')

# A self contained companion can be reviewed without a Word renderer.
parts=[]
for line in lines:
    if not line: continue
    if line.startswith('!['):
        caption=line.split(']')[0][2:]
        file=OUT/line.split('](')[1][:-1]
        encoded=base64.b64encode(file.read_bytes()).decode('ascii')
        cls='mobile' if 'mobile' in file.name else 'desktop'
        parts.append(f'<figure class="{cls}"><img alt="{html.escape(caption)}" src="data:image/png;base64,{encoded}"></figure>')
    elif line.startswith('#'):
        level=len(line)-len(line.lstrip('#'))
        parts.append(f'<h{level}>{html.escape(line[level:].strip())}</h{level}>')
    elif line.startswith('- [ ] '): parts.append('<p class="check">☐ '+html.escape(line[6:])+'</p>')
    elif line.startswith('- '): parts.append('<p class="bullet">• '+html.escape(line[2:])+'</p>')
    else: parts.append('<p>'+html.escape(line).replace('\n','<br>')+'</p>')
style='''body{margin:0;background:#f5f6f8;color:#192430;font:17px/1.65 system-ui,sans-serif}main{max-width:900px;margin:36px auto;padding:50px 64px;background:white}h1{font-size:36px;line-height:1.2}h2{font-size:27px;margin-top:52px;border-top:1px solid #dde2e8;padding-top:24px}h3{font-size:21px;margin-top:36px}p{margin:16px 0}figure{margin:32px 0 6px}img{width:100%;height:auto;border:1px solid #dde2e8}.mobile{max-width:340px;margin-left:auto;margin-right:auto}.check,.bullet{padding-left:8px} @media(max-width:650px){main{padding:24px;margin:0}body{font-size:16px}}@media print{body{background:white;font-size:11pt}main{margin:0;padding:0;max-width:none}h2{break-before:page}figure{break-inside:avoid}img{max-height:8in;object-fit:contain}.mobile{max-width:2.7in}}'''
(OUT/'course-creation-manual.html').write_text('<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>UniVybe course creation manual</title><style>'+style+'</style><main>'+''.join(parts)+'</main></html>', encoding='utf-8')

# Reader focused DOCX, with deliberate breaks before each training module.
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
normal = doc.styles['Normal']
normal.font.name='Calibri'; normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.08
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    doc.styles[name].font.color.rgb=RGBColor(0,0,0)
doc.styles['Title'].font.size=Pt(27)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(14)
header = section.header.paragraphs[0]
header.text='UNIVYBE  /  EDUCATOR TRAINING  /  LOCAL DRAFT'
header.style='Caption'; header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=section.footer.paragraphs[0]; footer.alignment=2
footer.add_run('Course creation manual  •  ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
skip_caption=False
for line in lines:
    if not line: continue
    if line.startswith('# '): doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        heading=line[3:]
        if heading.startswith('Module ') or heading=='Student portal walkthrough': doc.add_page_break()
        doc.add_heading(heading, 1)
    elif line.startswith('### '):
        # DOCX headings use words and spaces without outline punctuation.
        heading=line[4:]
        heading=heading.replace('.', ' ')
        doc.add_heading(heading, 2)
    elif line.startswith('!['):
        caption=line.split(']')[0][2:]
        file=line.split('](')[1][:-1]
        if file.endswith('01-student-course.png'): pass
        else: doc.add_page_break()
        img=OUT/file
        doc.add_picture(str(img), width=Inches(2.7 if 'mobile' in file else 6.7))
        p=doc.paragraphs[-1]; p.alignment=1; p.paragraph_format.keep_with_next=True
        doc.add_paragraph(caption+'. Illustrative data.', 'Caption')
        skip_caption=True
    elif skip_caption:
        skip_caption=False
    elif line.startswith('- [ ] '): doc.add_paragraph('☐ '+line[6:])
    elif line.startswith('- '): doc.add_paragraph(line[2:], 'List Bullet')
    else: doc.add_paragraph(line)
doc.save(OUT/'course-creation-manual.docx')
print(json.dumps(dict(modules=len(modules), lessons=course['estimated_lessons_count'], quizzes=course['estimated_quizzes_count'], pending_assets=len(assets))))
